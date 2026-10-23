#!/usr/bin/env python3
"""
Convert the Excel tracker into Backend/data/catalog.json.

Usage:
    pip install openpyxl
    python Backend/scripts/xlsx_to_json.py [path/to/workbook.xlsx]

The catalog holds the content that is the same for every user (roadmap tasks,
company rows, links...) plus the *default* values of the editable columns
(status, confidence, next action...). Each user's own edits are stored in the
database and layered on top.

Row ids are built from the row's text (e.g. "swe:dsa-big-o"), not its position,
so re-running this after editing the workbook keeps existing users' progress
attached to the right rows.

Some columns in the workbook have no header ("Column1", "Column2"...). They are
mapped here by their dropdowns and contents:
  Company Role Tracker  M = application status, N = applied date, O = notes, P = next action
  All Skills Reference  F = confidence, G = what to learn, H = practice task, I = resource, J = approach
"""
import datetime as dt
import json
import re
import sys
import warnings
from pathlib import Path

from openpyxl import load_workbook

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SOURCE = ROOT / "data" / "source" / "AI_ML_SWE_Integrated_Development_Tracker_Automated_Progress.xlsx"
SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE
OUT = ROOT / "data" / "catalog.json"

# ---------------------------------------------------------------- status sets
STATUS_SETS = {
    "learning": {
        "options": ["Not Started", "Pending", "Revision", "Done"],
        "initial": "Not Started", "done": "Done", "checkable": True,
        "tones": {"Not Started": "idle", "Pending": "progress", "Revision": "review", "Done": "done"},
    },
    "project": {
        "options": ["Not started", "In progress", "Needs revision", "Done"],
        "initial": "Not started", "done": "Done", "checkable": True,
        "tones": {"Not started": "idle", "In progress": "progress", "Needs revision": "review", "Done": "done"},
    },
    "application": {
        "options": ["Not started", "To research", "Ready to apply", "Applied", "Assessment",
                    "Interview", "Offer", "Rejected", "On hold"],
        "initial": "Not started", "done": "Offer", "checkable": False,
        "tones": {"Not started": "idle", "To research": "idle", "Ready to apply": "info", "Applied": "info",
                  "Assessment": "progress", "Interview": "progress", "Offer": "done",
                  "Rejected": "review", "On hold": "idle"},
    },
}
# Map old values into a set's options (the workbook mixes spellings).
ALIASES = {
    "project": {"not started": "Not started", "pending": "In progress", "revision": "Needs revision"},
    "learning": {"not started": "Not Started"},
    "application": {"not started": "Not started"},
}


def norm_status(value, set_name):
    spec = STATUS_SETS[set_name]
    if value in spec["options"]:
        return value
    mapped = ALIASES[set_name].get(str(value).strip().lower())
    return mapped if mapped in spec["options"] else spec["initial"]


# ---------------------------------------------------------------- column specs
def col(key, label, letter=None, type="text", table=False, editable=False):
    return {"key": key, "label": label, "letter": letter, "type": type, "table": table, "editable": editable}


def status_col(letter):
    return col("status", "Status", letter, "status", True, True)


def note_col(label="Notes", letter=None, key="notes"):
    return col(key, label, letter, "note", False, True)


ROADMAP_COLS = [
    col("phase", "Phase", "A"),
    col("category", "Category", "B"),
    col("skill", "Skill / task", "C"),
    col("priority", "Priority", "D", "badge", True),
    status_col("E"),
    col("confidence", "Confidence", "F", "confidence", True, True),
    col("learn", "What to learn / deliverable", "G", "long"),
    col("resource", "Resource / next action", "H", "auto"),
    note_col(),
]

ROADMAP_META = dict(
    kind="roadmap", group="learn", first_row=4, primary="skill", secondary="category",
    groupBy="phase", filters=["priority", "category"], statusSet="learning", idFrom=["skill"],
)

SHEETS = [
    dict(key="swe", sheet="SWE Roadmap", title="SWE Roadmap", columns=ROADMAP_COLS, **ROADMAP_META),
    dict(key="aiml", sheet="AI ML Roadmap", title="AI/ML Roadmap", columns=ROADMAP_COLS, **ROADMAP_META),
    dict(key="integration", sheet="AI + SWE Integration", title="AI + SWE Integration", columns=ROADMAP_COLS, **ROADMAP_META),
    dict(
        key="companies", sheet="Company Role Tracker", title="Company roles", kind="tracker", group="jobs",
        subtitle="Companies hiring for AI/ML and full-stack roles. Verify vacancies and eligibility before applying.",
        first_row=2, primary="company", secondary="track", groupBy=None,
        filters=["track", "priority"], statusSet="application", idFrom=["company", "roles"],
        help={"title": "What do the priorities mean?", "items": [
            {"term": "Priority", "text": "Reasonable first-wave targets to investigate."},
            {"term": "Stretch", "text": "More competitive, or role availability and eligibility may be harder."},
            {"term": "verify opening", "text": "Check that there is an active vacancy first."}]},
        columns=[
            col("company", "Company", "A"), col("track", "Track", "B"),
            col("roles", "Role titles to search", "C", "long"),
            col("priority", "Priority", "D", "badge", True),
            col("core", "Core skills required", "E", "long"),
            col("helpful", "Helpful / differentiating skills", "F", "long"),
            col("rounds", "Typical interview rounds (not guaranteed)", "G", "long"),
            col("formats", "Interview types / format", "H", "long"),
            col("prepare", "What to prepare / likely focus", "I", "long"),
            col("eligibility", "Eligibility note", "J", "long"),
            col("link", "Careers page", "K", "link", True),
            status_col("M"),  # header "Column2" in the workbook
            col("appliedDate", "Applied date", "N", "date", False, True),
            note_col("Notes", "O"),
            col("nextAction", "Next action", "P", "input", False, True),
        ],
    ),
    dict(
        key="apply", sheet="Live & Direct Apply Links", title="Apply links", kind="tracker", group="jobs",
        subtitle="Direct career pages and role links. A link isn't a guarantee that a vacancy is still open, "
                 "so check the deadline and eligibility before applying.",
        first_row=2, primary="company", secondary="roles", groupBy=None,
        filters=["sourceType"], statusSet="application", idFrom=["company", "roles"],
        columns=[
            col("company", "Company / portal", "A"), col("roles", "Role / track", "B"),
            col("link", "Apply link", "C", "link", True),
            col("checked", "Last checked", "D", "text", True),
            col("verify", "What to verify before applying", "E", "long"),
            col("sourceType", "Source type", "F"),
            status_col("G"),
            col("appliedDate", "Applied date", "H", "date", False, True),
            col("nextAction", "Next action", "I", "input", False, True),
            note_col(),
        ],
    ),
    dict(
        key="interview", sheet="Interview Round Guide", title="Interview rounds", kind="reference", group="jobs",
        subtitle="What each interview stage usually involves and how to prepare. Not every company uses every stage.",
        first_row=2, primary="stage", secondary=None, groupBy=None, filters=[], statusSet=None,
        idFrom=["stage"],
        columns=[
            col("stage", "Round / stage", "A"), col("formats", "Common formats", "B", "text", True),
            col("ai", "AI/ML focus", "C", "long"), col("fullstack", "Full Stack / React Native focus", "D", "long"),
            col("checklist", "Preparation checklist", "E", "long"), col("note", "Important note", "F", "long"),
        ],
    ),
    dict(
        key="skills", sheet="All Skills Reference", title="Skills reference", kind="catalog", group="library",
        subtitle="The broader skill catalogue. The roadmaps are the guided paths; use this list to look things up.",
        first_row=2, primary="skill", secondary="track", groupBy="category",
        filters=["track", "priority"], statusSet="learning", idFrom=["skill", "track"],
        columns=[
            col("track", "Track", "A"), col("category", "Category", "B"), col("skill", "Skill / topic", "C"),
            col("priority", "Priority", "D", "badge", True), status_col("E"),
            col("confidence", "Confidence", "F", "confidence", True, True),
            col("learn", "What to learn", "G", "long"), col("practice", "Practice task", "H", "long"),
            col("resource", "Resource", "I", "auto"), col("approach", "Approach", "J", "long"),
            note_col(),
        ],
    ),
    dict(
        key="resources", sheet="All Learning Resources", title="Learning resources", kind="catalog", group="library",
        subtitle="Documentation and courses behind the roadmaps.",
        first_row=2, primary="resource", secondary="domain", groupBy=None,
        filters=["domain"], statusSet="learning", idFrom=["resource"],
        columns=[
            col("domain", "Domain", "A"), col("resource", "Resource", "B"),
            col("link", "Link", "C", "link", True), col("useFor", "Use it for", "D", "text", True),
            status_col("E"), note_col("Notes", "F"),
        ],
    ),
    dict(
        key="projects", sheet="All Project Ideas", title="Project ideas", kind="catalog", group="library",
        subtitle="Milestones for portfolio projects, with space for proof links and target dates.",
        first_row=2, primary="item", secondary="milestone", groupBy="project",
        filters=["milestone"], statusSet="project", idFrom=["project", "item"],
        columns=[
            col("project", "Project", "A"), col("milestone", "Milestone", "B"),
            col("item", "Checklist item", "C"), status_col("D"),
            col("proof", "Proof / URL", "E", "input", False, True),
            col("nextAction", "Next action", "F", "input", False, True),
            col("targetDate", "Target date", "G", "date", False, True),
            note_col("Notes", "H"),
        ],
    ),
    dict(
        key="sources", sheet="Sources & Notes", title="Sources & notes", kind="reference", group="library",
        subtitle="Where the company and interview information came from. Processes change, so verify the current listing.",
        first_row=2, primary="source", secondary=None, groupBy=None, filters=[], statusSet=None,
        idFrom=["source"],
        columns=[
            col("source", "Source / company", "A"), col("supports", "What it supports", "B", "text", True),
            col("link", "Link", "C", "link", True), col("notes", "Notes", "D", "text", True),
        ],
    ),
]


# ---------------------------------------------------------------- helpers
def slug(text):
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:70]


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def letter_index(letter):
    n = 0
    for ch in letter:
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def cell_at(cells, letter):
    i = letter_index(letter)
    return cells[i] if i < len(cells) else None


def convert_sheet(wb, spec):
    ws = wb[spec["sheet"]]
    rows, seen, footnotes = [], {}, []
    in_footnotes = False
    primary_letter = next(c["letter"] for c in spec["columns"] if c["key"] == spec["primary"])

    for raw in ws.iter_rows(min_row=spec["first_row"], values_only=True):
        cells = list(raw)

        if spec["key"] == "sources":
            if not any(cell_text(c) for c in cells):
                in_footnotes = True  # a blank row separates the sources from the notes underneath
                continue
            if in_footnotes:
                footnotes.append({"label": cell_text(cells[0]), "text": cell_text(cells[1])})
                continue

        if not cell_text(cell_at(cells, primary_letter)):
            continue

        values = {}
        for c in spec["columns"]:
            raw_value = cell_at(cells, c["letter"]) if c["letter"] else None
            text = cell_text(raw_value)
            if c["type"] == "status":
                text = norm_status(text, spec["statusSet"])
            elif c["type"] == "confidence":
                ok = isinstance(raw_value, (int, float)) and 1 <= int(raw_value) <= 5
                text = int(raw_value) if ok else 1
            values[c["key"]] = text

        base = f'{spec["key"]}:{slug(" ".join(values[k] for k in spec["idFrom"]))}'
        seen[base] = seen.get(base, 0) + 1
        rows.append({"id": base if seen[base] == 1 else f"{base}-{seen[base]}", "values": values})

    subtitle = spec.get("subtitle") or cell_text(ws.cell(row=2, column=1).value)
    out = {k: spec[k] for k in ("key", "title", "kind", "group", "primary", "secondary", "groupBy", "filters", "statusSet")}
    out.update(
        subtitle=subtitle,
        help=spec.get("help"),
        footnotes=footnotes,
        columns=[{k: v for k, v in c.items() if k != "letter"} for c in spec["columns"]],
        rows=rows,
    )
    return out


def main():
    wb = load_workbook(SOURCE, data_only=True)
    sheets = [convert_sheet(wb, spec) for spec in SHEETS]
    OUT.write_text(json.dumps({"statusSets": STATUS_SETS, "sheets": sheets}, indent=1, ensure_ascii=False), encoding="utf-8")
    for s in sheets:
        print(f'{s["key"]:<12} {len(s["rows"]):>3} rows   {s["title"]}')
    print(f"\nWrote {OUT}")


if __name__ == "__main__":
    main()
